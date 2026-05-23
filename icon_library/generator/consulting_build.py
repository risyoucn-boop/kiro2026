"""Build script for the 200-icon consulting framework pack.

Emits:
    icon_library/icons_consulting/<CODE>.svg                              × 200
    icon_library/output/PPT图标语义库_咨询框架_200.pptx
    icon_library/output/PPT图标语义库_咨询框架_200_检索表.xlsx
    icon_library/output/PPT图标语义库_咨询框架_200_索引.csv

Independent of the 1000-icon general pack — different folder, different
filename prefix, different category codes. Run with `python consulting_build.py`.
"""
from __future__ import annotations

import csv
import os
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

sys.path.insert(0, os.path.dirname(__file__))

from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from primitives import normalize
from svg_writer import render_svg
from pptx_writer import add_icon_shape
from consulting_bases import CATEGORIES


# ---------------------------------------------------------------------------
# Output paths
# ---------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[1]
SVG_DIR = ROOT / "icons_consulting"
OUT_DIR = ROOT / "output"
SVG_DIR.mkdir(exist_ok=True, parents=True)
OUT_DIR.mkdir(exist_ok=True, parents=True)

PPTX_PATH = OUT_DIR / "PPT图标语义库_咨询框架_200.pptx"
XLSX_PATH = OUT_DIR / "PPT图标语义库_咨询框架_200_检索表.xlsx"
CSV_PATH  = OUT_DIR / "PPT图标语义库_咨询框架_200_索引.csv"


# ---------------------------------------------------------------------------
# Brand palette
# ---------------------------------------------------------------------------
ACCENT      = "#1F4E79"  # consulting-style deep blue
DARK_INK    = "#2D3748"
SOFT_GRAY   = "#A0AEC0"
TITLE_COLOR = "#1A202C"


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------
@dataclass
class IconDef:
    code: str
    cat_code: str
    cat_name: str
    cn_name: str
    en_name: str
    parts: list
    metaphor: str
    usage: str
    keywords_cn: list = field(default_factory=list)
    keywords_en: list = field(default_factory=list)
    license_note: str = ("原创设计 · 几何抽象 · 通用商业框架概念 · "
                         "非复制任何特定咨询机构视觉资产")


def build_all() -> list[IconDef]:
    out: list[IconDef] = []
    for cat_code, cat_name, base_list in CATEGORIES:
        idx = 1
        for entry in base_list:
            key, cn, en, fn, metaphor, usage, kw_cn, kw_en = entry
            out.append(IconDef(
                code=f"{cat_code}-{idx:03d}",
                cat_code=cat_code,
                cat_name=cat_name,
                cn_name=cn,
                en_name=en,
                parts=fn(),
                metaphor=metaphor,
                usage=usage,
                keywords_cn=list(kw_cn),
                keywords_en=list(kw_en),
            ))
            idx += 1
    return out


# ---------------------------------------------------------------------------
# Step 1: write all SVGs
# ---------------------------------------------------------------------------
def write_svgs(icons):
    print(f"[1/4] writing {len(icons)} SVG files -> {SVG_DIR}")
    t0 = time.time()
    for f in SVG_DIR.iterdir():
        if f.is_file() and f.suffix == ".svg":
            f.unlink()
    for ic in icons:
        norm = normalize(ic.parts)
        svg = render_svg(norm, fill=DARK_INK,
                         title=f"{ic.code} {ic.cn_name} / {ic.en_name}")
        (SVG_DIR / f"{ic.code}.svg").write_text(svg, encoding="utf-8")
    print(f"      done in {time.time()-t0:.1f}s")


# ---------------------------------------------------------------------------
# Step 2: master PPTX
# ---------------------------------------------------------------------------
ICONS_PER_ROW = 6        # fewer per row, bigger icons (frameworks need detail)
ROWS_PER_PAGE = 4

ICON_BOX = Inches(1.2)
COL_GAP  = Inches(0.6)
ROW_GAP  = Inches(0.7)
GRID_LEFT = Inches(0.7)
GRID_TOP  = Inches(1.5)


def write_pptx(icons):
    print(f"[2/4] writing master PPTX -> {PPTX_PATH}")
    t0 = time.time()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- intro slide ----
    s = prs.slides.add_slide(blank)
    _add_text(s, "PPT 图标语义库 · 咨询框架包 · 200 个", Inches(0.6), Inches(0.6),
              Inches(12), Inches(1.0), pt=32, bold=True, color=TITLE_COLOR)
    _add_text(s,
              "200 个原生 PowerPoint 形状 (custGeom)。"
              "每个图标对应一个公开的商业方法论框架（2x2矩阵、SWOT、五力、PDCA…）。"
              "所有设计为原创几何抽象，未复制任何特定咨询机构的视觉资产。",
              Inches(0.6), Inches(1.5), Inches(12), Inches(0.9),
              pt=12, color=DARK_INK)
    _add_text(s, "6 个分类 · 战略框架 / 流程模型 / 组织架构 / 数据可视化 / 诊断工具 / 战略概念",
              Inches(0.6), Inches(2.5), Inches(12), Inches(0.5),
              pt=13, color=SOFT_GRAY)

    swatch_x = Inches(0.6)
    swatch_y = Inches(3.5)
    sw = Inches(2.0)
    sh = Inches(1.7)
    for i, (code, name, base_list) in enumerate(CATEGORIES):
        col = i % 6
        row = i // 6
        x = swatch_x + (sw + Inches(0.1)) * col
        y = swatch_y + (sh + Inches(0.2)) * row
        first = next(ic for ic in icons if ic.cat_code == code)
        add_icon_shape(s, first.parts,
                       left=int(x + (sw - Inches(0.7)) / 2),
                       top=int(y + Inches(0.1)),
                       size=int(Inches(0.7)),
                       fill_hex=ACCENT,
                       name=f"category_{code}")
        _add_text(s, name, x, y + Inches(0.85), sw, Inches(0.4),
                  pt=10, bold=True, color=TITLE_COLOR, align=PP_ALIGN.CENTER)
        _add_text(s, f"{code} · {len(base_list)} icons",
                  x, y + Inches(1.20), sw, Inches(0.3),
                  pt=8, color=SOFT_GRAY, align=PP_ALIGN.CENTER)

    # ---- per-category cover + grid pages ----
    by_cat: dict[str, list] = {}
    for ic in icons:
        by_cat.setdefault(ic.cat_code, []).append(ic)

    for code, cat_name, _ in CATEGORIES:
        cat_icons = by_cat[code]
        cs = prs.slides.add_slide(blank)
        _add_text(cs, cat_name, Inches(0.6), Inches(2.6),
                  Inches(12), Inches(1.2), pt=42, bold=True, color=TITLE_COLOR)
        _add_text(cs, f"{code} · {len(cat_icons)} icons",
                  Inches(0.6), Inches(3.8), Inches(12), Inches(0.6),
                  pt=18, color=ACCENT)
        _add_text(cs, "每个图标对应一个公开商业方法论框架。Shape Fill 一键改色。",
                  Inches(0.6), Inches(4.6), Inches(12), Inches(0.6),
                  pt=12, color=SOFT_GRAY)

        per_page = ICONS_PER_ROW * ROWS_PER_PAGE
        for page_idx in range(0, len(cat_icons), per_page):
            page = cat_icons[page_idx: page_idx + per_page]
            slide = prs.slides.add_slide(blank)
            _add_text(slide, cat_name, Inches(0.6), Inches(0.3),
                      Inches(8), Inches(0.6), pt=20, bold=True, color=TITLE_COLOR)
            n_pages = (len(cat_icons) + per_page - 1) // per_page
            _add_text(slide,
                      f"{code} · 第 {page_idx // per_page + 1}/{n_pages} 页 · 共 {len(cat_icons)} 个",
                      Inches(0.6), Inches(0.85), Inches(10), Inches(0.4),
                      pt=10, color=SOFT_GRAY)

            for k, ic in enumerate(page):
                row = k // ICONS_PER_ROW
                col = k % ICONS_PER_ROW
                left = GRID_LEFT + (ICON_BOX + COL_GAP) * col
                top = GRID_TOP + (ICON_BOX + ROW_GAP) * row
                add_icon_shape(slide, ic.parts,
                               left=int(left), top=int(top),
                               size=int(ICON_BOX),
                               fill_hex=DARK_INK,
                               name=ic.code)
                _add_text(slide, ic.code,
                          left - Inches(0.2), top + ICON_BOX + Emu(20000),
                          ICON_BOX + Inches(0.4), Inches(0.22),
                          pt=8, bold=True, color=ACCENT,
                          align=PP_ALIGN.CENTER)
                _add_text(slide, ic.cn_name,
                          left - Inches(0.3), top + ICON_BOX + Inches(0.25),
                          ICON_BOX + Inches(0.6), Inches(0.22),
                          pt=8, color=DARK_INK, align=PP_ALIGN.CENTER)

    prs.save(str(PPTX_PATH))
    size_mb = PPTX_PATH.stat().st_size / 1_000_000
    print(f"      done in {time.time()-t0:.1f}s · {size_mb:.2f} MB · {len(prs.slides)} slides")


def _add_text(slide, text, left, top, width, height,
              pt=12, bold=False, color="#000000", align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(pt)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color.lstrip("#"))


# ---------------------------------------------------------------------------
# Step 3: Excel + CSV
# ---------------------------------------------------------------------------
EXCEL_HEADERS = [
    "编号", "图标预览", "中文名", "英文名", "图标来源", "图标集",
    "场景章节", "隐喻含义", "建议用法",
    "中文关键词", "英文关键词",
    "所在PPT页码", "所在PPT分类",
    "Office改色状态", "授权提示",
]


def write_excel(icons):
    print(f"[3/4] writing Excel + CSV -> {XLSX_PATH}")
    t0 = time.time()

    page_map: dict[str, int] = {}
    cur = 1
    by_cat: dict[str, list] = {}
    for ic in icons:
        by_cat.setdefault(ic.cat_code, []).append(ic)
    per_page = ICONS_PER_ROW * ROWS_PER_PAGE
    for code, _, _ in CATEGORIES:
        cur += 1
        cat_icons = by_cat[code]
        for page_idx in range(0, len(cat_icons), per_page):
            cur += 1
            for ic in cat_icons[page_idx: page_idx + per_page]:
                page_map[ic.code] = cur

    wb = Workbook()
    ws = wb.active
    ws.title = "咨询框架检索表"

    head_font = Font(bold=True, color="FFFFFF", size=11)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    border = Border(left=Side(style="thin", color="E2E8F0"),
                    right=Side(style="thin", color="E2E8F0"),
                    top=Side(style="thin", color="E2E8F0"),
                    bottom=Side(style="thin", color="E2E8F0"))
    for col_idx, h in enumerate(EXCEL_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_idx, ic in enumerate(icons, start=2):
        page_no = page_map.get(ic.code, "")
        row = [
            ic.code, "", ic.cn_name, ic.en_name,
            "原创设计·几何抽象", "PPT图标语义库·咨询框架包",
            ic.cat_name, ic.metaphor, ic.usage,
            " / ".join(ic.keywords_cn),
            ", ".join(ic.keywords_en),
            page_no, ic.cat_name,
            "已完成 (PowerPoint 原生 custGeom 形状, Shape Fill 直接改色)",
            ic.license_note,
        ]
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border

    widths = [12, 10, 24, 28, 22, 28, 18, 36, 36, 30, 30, 12, 18, 36, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    for r in range(2, len(icons) + 2):
        ws.row_dimensions[r].height = 24
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("使用说明与版权")
    instructions = [
        ["如何使用咨询框架包", ""],
        ["", ""],
        ["1. 检索", "在 Sheet1 用关键词筛/搜索（如 矩阵 / 漏斗 / S曲线 / 飞轮 / 拐点）"],
        ["2. 找到编号", "复制 [编号] 字段（例 STRAT-007 = 三层增长曲线）"],
        ["3. 找到 PPTX 页码", "查 [所在PPT页码] 字段"],
        ["4. 复制图标", "打开 PPT图标语义库_咨询框架_200.pptx，单击图标后 Ctrl+C"],
        ["5. 粘贴到业务 PPT", "在你的目标 PPT 里 Ctrl+V"],
        ["6. 一键改色", "选中图标 → 形状格式 → 形状填充 → 任意品牌色"],
        ["", ""],
        ["版权与设计声明", ""],
        ["原创性", "所有图标为原创几何抽象设计，使用基本几何形状（矩形/圆/多边形）"],
        ["概念来源", "图标对应的方法论概念（2x2矩阵/SWOT/五力/PDCA等）属于公开商业教育素材"],
        ["不复制", "未复制任何具体咨询机构（如 麦肯锡/BCG/贝恩/埃森哲 等）的视觉识别系统、配色、版式或专有图示"],
        ["可商用", "本资产可商用、可二次修改、可再分发"],
        ["免责", "如需引用某框架的方法论原文/详细解释，请自行查阅相应的公开商业教育资料或学术著作"],
        ["", ""],
        ["技术说明", ""],
        ["原生形状", "每个图标都是 PowerPoint 原生 custGeom freeform 形状，不是 SVG"],
        ["填充模型", "单色 Solid Fill。所有路径同色，Shape Fill 一次改色全部生效"],
        ["WPS 兼容", "custGeom 是 OOXML 标准，WPS / Office 365 / Office 2016+ / Keynote 都能识别"],
    ]
    for r, (a, b) in enumerate(instructions, 1):
        ws2.cell(row=r, column=1, value=a).font = Font(bold=True, size=11)
        ws2.cell(row=r, column=2, value=b)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 95

    wb.save(XLSX_PATH)

    with open(CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(EXCEL_HEADERS)
        for ic in icons:
            page_no = page_map.get(ic.code, "")
            w.writerow([
                ic.code, "", ic.cn_name, ic.en_name,
                "原创设计·几何抽象", "PPT图标语义库·咨询框架包",
                ic.cat_name, ic.metaphor, ic.usage,
                " / ".join(ic.keywords_cn),
                ", ".join(ic.keywords_en),
                page_no, ic.cat_name,
                "已完成 (PowerPoint 原生 custGeom 形状, Shape Fill 直接改色)",
                ic.license_note,
            ])
    print(f"      done in {time.time()-t0:.1f}s")


# ---------------------------------------------------------------------------
def main():
    print(">> Building PPT 图标语义库 · 咨询框架包 · 200 个")
    t0 = time.time()
    icons = build_all()
    print(f"   loaded {len(icons)} icons across {len(CATEGORIES)} categories")

    write_svgs(icons)
    write_pptx(icons)
    write_excel(icons)

    print(f"\n[4/4] DONE in {time.time()-t0:.1f}s")
    print(f"      SVG       : {SVG_DIR}/  ({len(icons)} files)")
    print(f"      PPTX      : {PPTX_PATH}  ({PPTX_PATH.stat().st_size/1_000_000:.2f} MB)")
    print(f"      Excel     : {XLSX_PATH}  ({XLSX_PATH.stat().st_size/1_000:.1f} KB)")
    print(f"      CSV       : {CSV_PATH}  ({CSV_PATH.stat().st_size/1_000:.1f} KB)")


if __name__ == "__main__":
    main()
