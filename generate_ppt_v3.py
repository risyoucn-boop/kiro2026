# -*- coding: utf-8 -*-
"""
银卡客服岗位 1+N 数字化转型 · 集团董事长汇报版（6 页）
- 平安橙 + 白底
- 战略级语言 / 删除人名 / 删除 12 项目清单 / 抽象为 3 类岗位能力
- 每页 1 个主 exhibit + 留白克制
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ===== 平安橙 (集团级版本：色彩更克制) =====
PA_ORANGE      = RGBColor(0xEC, 0x6F, 0x1C)
PA_ORANGE_DARK = RGBColor(0xC8, 0x55, 0x0E)
PA_ORANGE_DEEP = RGBColor(0x7A, 0x32, 0x06)
PA_ORANGE_LT   = RGBColor(0xFF, 0xE8, 0xD2)
PA_ORANGE_BG   = RGBColor(0xFF, 0xF7, 0xEF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x0C, 0x0C, 0x0C)
DARK      = RGBColor(0x2A, 0x2A, 0x2A)
GRAY      = RGBColor(0x66, 0x66, 0x66)
MIDGRAY   = RGBColor(0x9E, 0x9E, 0x9E)
LIGHTGRAY = RGBColor(0xD8, 0xD8, 0xD8)
PALE      = RGBColor(0xFA, 0xFA, 0xFA)

CN_FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height


def set_run(run, text, size=11, bold=False, color=INK, italic=False):
    run.text = text
    run.font.name = CN_FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', CN_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def tb(slide, l, t, w, h, text, size=11, bold=False, color=INK,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, line_spacing=1.2):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        set_run(run, ln, size=size, bold=bold, color=color, italic=italic)
    return box


def rect(slide, l, t, w, h, fill, line=None, lw=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def chevron(slide, l, t, w, h, fill, text, text_size=11, text_color=WHITE):
    """向右的箭头形 chevron — 用于复制路径图"""
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    set_run(run, text, size=text_size, bold=True, color=text_color)
    return s


def white_bg(slide):
    rect(slide, 0, 0, SW, SH, WHITE)


def board_header(slide, page, total, eyebrow, action_title, insight=None):
    """
    集团董事长版页眉：极简、克制、留白。
    - 顶部细橙色条
    - 左上：项目名（极小灰色）；右上：页码
    - 主标题：超大字，黑色
    - 副线：如有
    - 底部一根橙色短粗 + 灰色细长分割线
    """
    rect(slide, 0, 0, SW, Inches(0.06), PA_ORANGE)

    # 左上
    tb(slide, Inches(0.5), Inches(0.18), Inches(8.0), Inches(0.25),
       eyebrow, size=9, bold=True, color=GRAY)
    # 右上
    tb(slide, Inches(11.4), Inches(0.18), Inches(1.4), Inches(0.25),
       f"{page} / {total}", size=9, bold=True, color=PA_ORANGE,
       align=PP_ALIGN.RIGHT)

    # 主标题
    tb(slide, Inches(0.5), Inches(0.55), Inches(12.3), Inches(0.85),
       action_title, size=24, bold=True, color=INK,
       anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)

    if insight:
        tb(slide, Inches(0.5), Inches(1.42), Inches(12.3), Inches(0.42),
           insight, size=13, color=PA_ORANGE_DARK, line_spacing=1.25,
           anchor=MSO_ANCHOR.MIDDLE)
        rule_y = Inches(1.92)
    else:
        rule_y = Inches(1.55)

    # 分割线（橙色短粗 + 灰色细长）
    rect(slide, Inches(0.5), rule_y, Inches(0.5), Inches(0.04), PA_ORANGE)
    rect(slide, Inches(1.0), rule_y + Inches(0.014), Inches(11.83),
         Inches(0.012), LIGHTGRAY)


def board_footer(slide):
    rect(slide, Inches(0.5), Inches(7.18), Inches(12.33), Inches(0.012), LIGHTGRAY)
    tb(slide, Inches(0.5), Inches(7.22), Inches(8.0), Inches(0.22),
       "资料来源：客服中心银卡试点项目组   ·   数字均为估算口径，待业管/财务校准",
       size=8.5, color=MIDGRAY, italic=True)
    tb(slide, Inches(8.5), Inches(7.22), Inches(4.33), Inches(0.22),
       "PING AN GROUP    ·    BOARD-LEVEL    ·    CONFIDENTIAL",
       size=8.5, bold=True, color=PA_ORANGE,
       align=PP_ALIGN.RIGHT)


def board_kicker(slide, l, t, w, h, label, body, fill=PA_ORANGE):
    rect(slide, l, t, w, h, fill)
    tb(slide, l + Inches(0.25), t + Inches(0.08),
       Inches(2.0), Inches(0.3),
       label, size=10, bold=True, color=PA_ORANGE_LT)
    tb(slide, l + Inches(0.25), t + Inches(0.38),
       w - Inches(0.5), h - Inches(0.45),
       body, size=14, bold=True, color=WHITE, line_spacing=1.3)


# ====================================================================
# Slide 1 — 战略定位
# ====================================================================
def slide_1():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)

    # 极简封面式页头
    rect(s, 0, 0, SW, Inches(0.08), PA_ORANGE)
    rect(s, 0, Inches(0.08), Inches(0.5), SH - Inches(0.08), PA_ORANGE_BG)

    tb(s, Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.3),
       "向集团董事长汇报   ·   银卡客服试点", size=10, bold=True, color=GRAY)
    tb(s, Inches(11.4), Inches(0.4), Inches(1.4), Inches(0.3),
       "1 / 6", size=10, bold=True, color=PA_ORANGE, align=PP_ALIGN.RIGHT)

    # 战略定位 eyebrow
    tb(s, Inches(0.8), Inches(1.05), Inches(11.5), Inches(0.4),
       "STRATEGIC POSITIONING   ·   战略定位",
       size=11, bold=True, color=PA_ORANGE_DARK)

    # 主标题（超大、黑、压顶）
    tb(s, Inches(0.8), Inches(1.45), Inches(11.7), Inches(1.25),
       "客服岗位 AI 化",
       size=46, bold=True, color=INK, line_spacing=1.0)
    tb(s, Inches(0.8), Inches(2.55), Inches(11.7), Inches(0.85),
       "—— 平安「一个客户、多个产品」战略在前端的最后一公里",
       size=22, bold=True, color=PA_ORANGE_DARK, line_spacing=1.2)

    # 橙色短粗分割
    rect(s, Inches(0.8), Inches(3.55), Inches(0.6), Inches(0.06), PA_ORANGE)

    # 三个战略锚 — 极简卡片（无装饰，仅左侧色块）
    anchors = [
        ("01",
         "综合金融",
         "单客户多产品需要前端岗位具备\n全场景串联与价值经营能力"),
        ("02",
         "AI 战略",
         "客服触点是集团 AI 落地\n最大规模、最易显效的前端场景"),
        ("03",
         "平安 One",
         "客户经营升级需要前端岗位\n从「服务承接」走向「价值经营」"),
    ]
    a_top = Inches(3.85)
    a_w = Inches(3.95)
    a_h = Inches(2.05)
    a_gap = Inches(0.1)
    for i, (no, name, desc) in enumerate(anchors):
        x = Inches(0.8) + (a_w + a_gap) * i
        rect(s, x, a_top, Inches(0.06), a_h, PA_ORANGE)
        tb(s, x + Inches(0.25), a_top + Inches(0.1),
           Inches(2.0), Inches(0.4),
           no, size=11, bold=True, color=PA_ORANGE)
        tb(s, x + Inches(0.25), a_top + Inches(0.4),
           a_w - Inches(0.4), Inches(0.6),
           name, size=24, bold=True, color=INK)
        tb(s, x + Inches(0.25), a_top + Inches(1.05),
           a_w - Inches(0.4), Inches(1.0),
           desc, size=12, color=DARK, line_spacing=1.4)

    # 底部主张条
    board_kicker(
        s, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.85),
        "CORE PROPOSITION   ·   核心主张",
        "客户在哪里，岗位 AI 化就跟到哪里 — 银卡作为试验场，三类岗位能力先跑通，再向集团客户经营体系全面复制。",
    )

    tb(s, Inches(0.8), Inches(7.2), Inches(11.7), Inches(0.22),
       "PING AN GROUP    ·    BOARD-LEVEL    ·    CONFIDENTIAL",
       size=8.5, bold=True, color=PA_ORANGE, align=PP_ALIGN.RIGHT)


# ====================================================================
# Slide 2 — 核心判断
# ====================================================================
def slide_2():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    board_header(
        s, 2, 6,
        "向集团董事长汇报   ·   核心判断",
        "客服中心已进入「AI + 岗位」重构期：这是集团 AI 战略在前端最大规模、最易显效的落地场景",
        "三大结构性压力同时到达 — 仅靠加人、加培训、加流程已难以承接，必须用 AI 重构岗位能力底座",
    )

    # 上半：三大结构性压力 (3 列)
    pressures = [
        ("服务复杂度",
         "客户结构持续高端化、监管强度持续上升、\n投诉与升级风险高度依赖少数资深主管经验",
         "+",
         "复杂客诉年增长率",
         "≈ 18%",
         "估算口径"),
        ("经营深度",
         "服销、私财、高客、挽留、顶私、平安 One 经营\n任务叠加 — 一线执行动作不稳定",
         "×",
         "前端经营任务叠加倍数",
         "≈ 3.2 倍",
         "估算口径"),
        ("人力天花板",
         "客户规模增速 > 人员增速 — \n靠「加人 / 加班」模式不可持续",
         "÷",
         "客户/坐席比 五年涨幅",
         "≈ 1.6 倍",
         "估算口径"),
    ]
    p_top = Inches(2.1)
    p_w = Inches(4.05)
    p_h = Inches(2.6)
    p_gap = Inches(0.09)
    for i, (k, desc, sym, m_lbl, m_val, m_note) in enumerate(pressures):
        x = Inches(0.5) + (p_w + p_gap) * i
        rect(s, x, p_top, p_w, p_h, WHITE, line=LIGHTGRAY, lw=0.75)
        # 顶部色条
        rect(s, x, p_top, p_w, Inches(0.08), PA_ORANGE)
        # 大符号
        tb(s, x + Inches(0.2), p_top + Inches(0.2),
           Inches(0.7), Inches(0.7),
           sym, size=36, bold=True, color=PA_ORANGE_LT,
           anchor=MSO_ANCHOR.MIDDLE)
        # 主名
        tb(s, x + Inches(0.95), p_top + Inches(0.25),
           p_w - Inches(1.1), Inches(0.5),
           k, size=20, bold=True, color=INK)
        # 描述
        tb(s, x + Inches(0.25), p_top + Inches(1.0),
           p_w - Inches(0.5), Inches(0.95),
           desc, size=11, color=DARK, line_spacing=1.4)
        # 度量
        rect(s, x + Inches(0.25), p_top + Inches(1.95),
             p_w - Inches(0.5), Inches(0.5), PA_ORANGE_BG)
        tb(s, x + Inches(0.4), p_top + Inches(1.97),
           p_w - Inches(2.0), Inches(0.5),
           m_lbl, size=10, bold=True, color=GRAY,
           anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + p_w - Inches(1.65), p_top + Inches(1.97),
           Inches(1.4), Inches(0.5),
           m_val, size=18, bold=True, color=PA_ORANGE_DARK,
           align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # 下半：核心判断 (4 个数字 — 量级佐证)
    rect(s, Inches(0.5), Inches(4.95), Inches(12.33), Inches(0.4), PA_ORANGE_BG)
    tb(s, Inches(0.7), Inches(4.95), Inches(12.0), Inches(0.4),
       "▍ 量级佐证 — 客服触点是集团 AI 落地最大规模场景",
       size=12, bold=True, color=PA_ORANGE_DARK, anchor=MSO_ANCHOR.MIDDLE)

    metrics = [
        ("≈ 25,000+",  "集团客服坐席规模",      "前端 AI 可触达岗位"),
        ("≈ 2.4 亿",   "集团个人客户基数",       "AI 化客户经营覆盖面"),
        ("≈ 7 亿+/年", "客服触点交互次数",       "AI 训练与显效高频场景"),
        ("≈ 35 亿+",   "客服年人力成本估算",     "AI 化潜在效率释放空间"),
    ]
    m_top = Inches(5.45)
    m_w = Inches(3.04)
    m_h = Inches(1.55)
    for i, (val, k, sub) in enumerate(metrics):
        x = Inches(0.5) + (m_w + Inches(0.04)) * i
        rect(s, x, m_top, m_w, m_h, WHITE, line=LIGHTGRAY, lw=0.75)
        rect(s, x, m_top, Inches(0.06), m_h, PA_ORANGE)
        tb(s, x + Inches(0.2), m_top + Inches(0.15),
           m_w - Inches(0.3), Inches(0.65),
           val, size=24, bold=True, color=PA_ORANGE_DARK)
        tb(s, x + Inches(0.2), m_top + Inches(0.85),
           m_w - Inches(0.3), Inches(0.3),
           k, size=11, bold=True, color=INK)
        tb(s, x + Inches(0.2), m_top + Inches(1.18),
           m_w - Inches(0.3), Inches(0.3),
           sub, size=9, color=GRAY)

    board_footer(s)


# ====================================================================
# Slide 3 — 银卡试验场 / 三类岗位能力
# ====================================================================
def slide_3():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    board_header(
        s, 3, 6,
        "向集团董事长汇报   ·   银卡试验场",
        "银卡先行 — 把客服岗位抽象为三类能力，先在最难的高端客服场景跑通，再向集团复制",
        "不再讲「12 个项目」 — 讲「3 类岗位能力底座」：服务化解 · 经营触达 · 知识生产",
    )

    # 三类岗位能力 — 大三柱
    capabilities = [
        {
            "code": "Ⅰ",
            "name": "服务化解能力",
            "en": "SERVICE RESOLUTION",
            "core": "把资深主管的隐性投诉 / 风险化解经验，沉淀为组织化、可训练、可追踪的能力",
            "scenes": "投诉化解  ·  服务现管  ·  消保前置  ·  舆情风险识别",
            "kpi_label": "资深主管重复兜底时长",
            "kpi_val": "↓ 40%",
            "kpi_note": "试点目标值",
        },
        {
            "code": "Ⅱ",
            "name": "经营触达能力",
            "en": "VALUE ENGAGEMENT",
            "core": "把单客户多产品的经营动作，从「人工盯盘」升级为「数据识别 + 场景化触达」",
            "scenes": "服销现管  ·  高客经营  ·  挽留转化  ·  顶私 / 私财场景",
            "kpi_label": "尾部识别 + 辅导闭环效率",
            "kpi_val": "↑ 3 倍",
            "kpi_note": "试点目标值",
        },
        {
            "code": "Ⅲ",
            "name": "知识生产能力",
            "en": "KNOWLEDGE PRODUCTION",
            "core": "把人工知识拆解、培训、运维链路，升级为 AI 辅助生产 + 标准模板的工业化体系",
            "scenes": "智能客服知识库  ·  新人育成  ·  教练辅导  ·  AI 陪练",
            "kpi_label": "知识拆解 / 上线周期",
            "kpi_val": "↓ 50%+",
            "kpi_note": "试点目标值",
        },
    ]
    c_top = Inches(2.1)
    c_w = Inches(4.05)
    c_h = Inches(4.0)
    c_gap = Inches(0.1)
    for i, c in enumerate(capabilities):
        x = Inches(0.5) + (c_w + c_gap) * i
        # 卡片
        rect(s, x, c_top, c_w, c_h, WHITE, line=LIGHTGRAY, lw=0.75)
        # 顶部色块
        rect(s, x, c_top, c_w, Inches(0.85), PA_ORANGE)
        tb(s, x + Inches(0.2), c_top + Inches(0.05),
           Inches(0.6), Inches(0.4),
           c["code"], size=24, bold=True, color=WHITE,
           italic=True, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + Inches(0.2), c_top + Inches(0.4),
           c_w - Inches(0.4), Inches(0.4),
           c["en"], size=9, bold=True, color=PA_ORANGE_LT)

        # 主名
        tb(s, x + Inches(0.25), c_top + Inches(1.0),
           c_w - Inches(0.5), Inches(0.55),
           c["name"], size=22, bold=True, color=INK)

        # 短分割
        rect(s, x + Inches(0.25), c_top + Inches(1.6),
             Inches(0.5), Inches(0.04), PA_ORANGE)

        # 核心定义
        tb(s, x + Inches(0.25), c_top + Inches(1.75),
           c_w - Inches(0.5), Inches(0.85),
           c["core"], size=11, color=DARK, line_spacing=1.5)

        # 覆盖场景
        tb(s, x + Inches(0.25), c_top + Inches(2.65),
           c_w - Inches(0.5), Inches(0.25),
           "覆盖场景", size=9, bold=True, color=GRAY)
        tb(s, x + Inches(0.25), c_top + Inches(2.88),
           c_w - Inches(0.5), Inches(0.55),
           c["scenes"], size=10.5, bold=True, color=PA_ORANGE_DARK,
           line_spacing=1.4)

        # KPI 块
        rect(s, x + Inches(0.25), c_top + Inches(3.45),
             c_w - Inches(0.5), Inches(0.5), PA_ORANGE_BG)
        tb(s, x + Inches(0.4), c_top + Inches(3.46),
           c_w - Inches(2.5), Inches(0.5),
           c["kpi_label"], size=9.5, color=GRAY,
           anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + c_w - Inches(2.05), c_top + Inches(3.46),
           Inches(1.8), Inches(0.5),
           c["kpi_val"], size=18, bold=True, color=PA_ORANGE_DARK,
           align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # 底部主张
    board_kicker(
        s, Inches(0.5), Inches(6.25), Inches(12.33), Inches(0.85),
        "BANK CARD AS LAB   ·   银卡试验场",
        "三类能力先在银卡跑通 — 形成「岗位 1+N」标准方法论，再向集团客服全量及综合金融客户经营体系复制。",
    )
    board_footer(s)


# ====================================================================
# Slide 4 — 规模与复制路径
# ====================================================================
def slide_4():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    board_header(
        s, 4, 6,
        "向集团董事长汇报   ·   规模与复制路径",
        "复制路径：银卡试点 → 客服中心全量 → 综合金融客户经营体系 — 三阶段、五年规模化",
        "复制不是「方法论 PPT」 — 是带着标准能力底座、标准 KPI、标准训练材料一站到底",
    )

    # 上半：复制路径 chevron 链
    path_top = Inches(2.15)
    path_h = Inches(1.5)
    stages = [
        {
            "title": "STAGE 1   银卡试点",
            "year": "Y1",
            "scope": "银卡客服 ≈ 3,500 坐席",
            "deliver": "三类能力跑通\n标准方法论沉淀",
            "color": PA_ORANGE,
        },
        {
            "title": "STAGE 2   客服中心全量",
            "year": "Y2 ~ Y3",
            "scope": "信用卡 + 寿险 + 产险 + 健康险\n客服 ≈ 25,000 坐席",
            "deliver": "三类能力规模化部署\n人均产能 ↑ · 单坐席成本 ↓",
            "color": PA_ORANGE_DARK,
        },
        {
            "title": "STAGE 3   集团客户经营",
            "year": "Y4 ~ Y5",
            "scope": "综合金融 2.4 亿个人客户\n+ 平安 One 客户经营全场景",
            "deliver": "前端 AI 化标杆\n客户价值经营全面升级",
            "color": PA_ORANGE_DEEP,
        },
    ]
    seg_w = Inches(4.05)
    seg_gap = Inches(0.06)
    for i, st in enumerate(stages):
        x = Inches(0.5) + (seg_w + seg_gap) * i
        # chevron 顶条
        chevron(s, x, path_top, seg_w, Inches(0.5),
                st["color"], st["title"], text_size=11.5, text_color=WHITE)
        # 卡片底
        rect(s, x, path_top + Inches(0.55), seg_w, path_h - Inches(0.5),
             WHITE, line=LIGHTGRAY, lw=0.75)
        rect(s, x, path_top + Inches(0.55), Inches(0.06), path_h - Inches(0.5),
             st["color"])
        tb(s, x + Inches(0.2), path_top + Inches(0.62),
           Inches(1.0), Inches(0.3),
           st["year"], size=11, bold=True, color=st["color"])
        tb(s, x + Inches(0.2), path_top + Inches(0.92),
           seg_w - Inches(0.4), Inches(0.5),
           st["scope"], size=11, bold=True, color=INK, line_spacing=1.3)
        tb(s, x + Inches(0.2), path_top + Inches(1.45),
           seg_w - Inches(0.4), Inches(0.6),
           st["deliver"], size=10, color=DARK, line_spacing=1.4)

    # 中部分割
    rect(s, Inches(0.5), Inches(4.05), Inches(12.33), Inches(0.04), PA_ORANGE_LT)
    tb(s, Inches(0.5), Inches(4.12), Inches(12.33), Inches(0.3),
       "▍ 价值规模 — 五年视角下，三类岗位能力 AI 化对集团释放的潜在价值",
       size=11, bold=True, color=PA_ORANGE_DARK)

    # 下半：4 个规模数字 — 大字
    values = [
        ("25,000+",  "可复制坐席",       "覆盖集团客服全量"),
        ("≈ 30%",    "前端人力释放",     "潜在节约工时占比 · 估算"),
        ("≈ 10 亿+", "年化效率价值",     "等价人力释放 · 估算"),
        ("2.4 亿",   "服务的客户基数",   "综合金融个人客户全覆盖"),
    ]
    v_top = Inches(4.5)
    v_w = Inches(3.04)
    v_h = Inches(1.65)
    for i, (val, k, sub) in enumerate(values):
        x = Inches(0.5) + (v_w + Inches(0.04)) * i
        rect(s, x, v_top, v_w, v_h, PA_ORANGE_BG)
        rect(s, x, v_top, v_w, Inches(0.06), PA_ORANGE)
        tb(s, x + Inches(0.2), v_top + Inches(0.2),
           v_w - Inches(0.3), Inches(0.75),
           val, size=30, bold=True, color=PA_ORANGE_DARK)
        tb(s, x + Inches(0.2), v_top + Inches(1.0),
           v_w - Inches(0.3), Inches(0.3),
           k, size=12, bold=True, color=INK)
        tb(s, x + Inches(0.2), v_top + Inches(1.32),
           v_w - Inches(0.3), Inches(0.3),
           sub, size=9, color=GRAY)

    # 底部主张
    board_kicker(
        s, Inches(0.5), Inches(6.3), Inches(12.33), Inches(0.8),
        "REPLICATION TAKEAWAY   ·   复制主张",
        "银卡 → 客服 → 集团 — 不是项目堆叠，而是一套岗位能力底座沿着「客户在哪里、AI 跟到哪里」复制。",
    )
    board_footer(s)


# ====================================================================
# Slide 5 — 机制保障 / 集团层面四项支持
# ====================================================================
def slide_5():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    board_header(
        s, 5, 6,
        "向集团董事长汇报   ·   机制保障",
        "请求集团四项支持 — 把试点变成集团样板，需要从机制上锚定",
        "不要工具、不要预算、不要新增编制 — 要授权、要数据、要协同、要激励",
    )

    asks = [
        {
            "no": "01",
            "name": "试点授权",
            "en": "PILOT MANDATE",
            "ask": "授予银卡作为「集团客服岗位 AI 化试点单元」的明确身份",
            "why": "有「集团身份」才能跨条线协同 — 否则只是子单位内部项目",
            "decide": "集团董事长授权 · 集团数字化委员会备案",
        },
        {
            "no": "02",
            "name": "数据合规",
            "en": "DATA COMPLIANCE",
            "ask": "建立客户数据脱敏可用机制 — 覆盖客服、投诉、高客、平安 One 场景",
            "why": "AI 化的核心要素是真实场景数据 — 没有合规数据，能力底座做不实",
            "decide": "集团合规 + 数据条线 + 安全条线三方联动",
        },
        {
            "no": "03",
            "name": "跨条线协同",
            "en": "CROSS-LINE SYNERGY",
            "ask": "明确客服中心 ↔ 业管 ↔ 智能客服 ↔ 银行端 ↔ 总部数字化协同口径",
            "why": "岗位 AI 化不是单一团队能完成 — 必须从机制上把它装进集团协同体系",
            "decide": "集团数字化条线统筹 + 季度联席会议",
        },
        {
            "no": "04",
            "name": "评价激励",
            "en": "RECOGNITION & INCENTIVE",
            "ask": "把岗位 AI 化产出纳入年度评价、数字化创新、绩效与人才发展通道",
            "why": "改革要靠骨干 — 没有「这是我的成绩」 · 一线骨干不会真干",
            "decide": "集团人力 + 数字化条线 联合机制",
        },
    ]
    # 2 x 2
    a_top = Inches(2.1)
    a_w = Inches(6.16)
    a_h = Inches(2.05)
    a_gap_x = Inches(0.01)
    a_gap_y = Inches(0.13)
    for i, a in enumerate(asks):
        col = i % 2
        row = i // 2
        x = Inches(0.5) + (a_w + a_gap_x) * col
        y = a_top + (a_h + a_gap_y) * row
        rect(s, x, y, a_w, a_h, WHITE, line=LIGHTGRAY, lw=0.75)
        # 编号块
        rect(s, x, y, Inches(1.1), a_h, PA_ORANGE)
        tb(s, x, y + Inches(0.2), Inches(1.1), Inches(0.55),
           a["no"], size=28, bold=True, color=WHITE,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x, y + Inches(0.85), Inches(1.1), Inches(0.3),
           a["en"], size=8, bold=True, color=PA_ORANGE_LT,
           align=PP_ALIGN.CENTER)
        # 主名
        tb(s, x + Inches(1.25), y + Inches(0.12),
           a_w - Inches(1.4), Inches(0.45),
           a["name"], size=18, bold=True, color=INK)
        # 诉求
        tb(s, x + Inches(1.25), y + Inches(0.6),
           a_w - Inches(1.4), Inches(0.5),
           a["ask"], size=11, bold=True, color=PA_ORANGE_DARK,
           line_spacing=1.35)
        # 为什么
        tb(s, x + Inches(1.25), y + Inches(1.15),
           Inches(2.3), Inches(0.25),
           "WHY", size=8.5, bold=True, color=GRAY)
        tb(s, x + Inches(1.25), y + Inches(1.38),
           a_w - Inches(1.4), Inches(0.4),
           a["why"], size=10, color=DARK, line_spacing=1.35)
        # decide
        rect(s, x + Inches(1.25), y + Inches(1.74),
             a_w - Inches(1.4), Inches(0.25), PA_ORANGE_BG)
        tb(s, x + Inches(1.35), y + Inches(1.74),
           a_w - Inches(1.6), Inches(0.25),
           f"决策口径   {a['decide']}", size=9, bold=True,
           color=PA_ORANGE_DARK, anchor=MSO_ANCHOR.MIDDLE)

    # 底部主张
    board_kicker(
        s, Inches(0.5), Inches(6.45), Inches(12.33), Inches(0.7),
        "GOVERNANCE TAKEAWAY   ·   治理主张",
        "四项支持落到位，银卡试点才能在 12 个月内成为集团样板 — 否则只是又一个内部项目。",
    )
    board_footer(s)


# ====================================================================
# Slide 6 — 愿景 / 三阶段目标
# ====================================================================
def slide_6():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    board_header(
        s, 6, 6,
        "向集团董事长汇报   ·   愿景",
        "一年样板 · 三年标杆 · 五年范式 — 客服岗位 AI 化重塑平安客户经营前端",
    )

    # 三阶段大型时间轴
    rect(s, Inches(0.5), Inches(2.15), Inches(12.33), Inches(0.06), PA_ORANGE_LT)

    visions = [
        {
            "horizon": "1 YEAR",
            "year_lbl": "一年内",
            "title": "集团客服 AI 化样板",
            "outcome": "三类岗位能力底座 ✓\n标准 KPI 与方法论 ✓\n首批可申报 / 可复制案例 ✓",
            "scope": "银卡 ≈ 3,500 坐席",
            "color": PA_ORANGE,
        },
        {
            "horizon": "3 YEARS",
            "year_lbl": "三年内",
            "title": "综合金融客户经营 AI 化标杆",
            "outcome": "客服全量规模化部署 ✓\n人均产能与服务承载显著提升 ✓\n岗位 AI 化方法论沉淀为集团资产 ✓",
            "scope": "客服 ≈ 25,000 坐席",
            "color": PA_ORANGE_DARK,
        },
        {
            "horizon": "5 YEARS",
            "year_lbl": "五年内",
            "title": "平安客户经营 AI 化范式",
            "outcome": "前端岗位与 AI 形成稳定协作模式 ✓\n客户价值经营前端化、智能化、个性化 ✓\n成为综合金融行业标杆 ✓",
            "scope": "服务 2.4 亿 个人客户",
            "color": PA_ORANGE_DEEP,
        },
    ]
    v_top = Inches(2.4)
    v_w = Inches(4.05)
    v_h = Inches(3.95)
    v_gap = Inches(0.1)
    for i, v in enumerate(visions):
        x = Inches(0.5) + (v_w + v_gap) * i
        # 卡片
        rect(s, x, v_top, v_w, v_h, WHITE, line=LIGHTGRAY, lw=0.75)
        # 顶部时间块
        rect(s, x, v_top, v_w, Inches(0.95), v["color"])
        tb(s, x + Inches(0.25), v_top + Inches(0.1),
           v_w - Inches(0.5), Inches(0.35),
           v["horizon"], size=10, bold=True, color=PA_ORANGE_LT)
        tb(s, x + Inches(0.25), v_top + Inches(0.4),
           v_w - Inches(0.5), Inches(0.55),
           v["year_lbl"], size=24, bold=True, color=WHITE)

        # 主目标
        tb(s, x + Inches(0.25), v_top + Inches(1.15),
           v_w - Inches(0.5), Inches(0.85),
           v["title"], size=18, bold=True, color=INK, line_spacing=1.2)

        # 短分割
        rect(s, x + Inches(0.25), v_top + Inches(2.05),
             Inches(0.5), Inches(0.04), v["color"])

        # outcome
        tb(s, x + Inches(0.25), v_top + Inches(2.2),
           v_w - Inches(0.5), Inches(1.3),
           v["outcome"], size=10.5, color=DARK, line_spacing=1.6)

        # 规模标签
        rect(s, x + Inches(0.25), v_top + Inches(3.4),
             v_w - Inches(0.5), Inches(0.4), PA_ORANGE_BG)
        tb(s, x + Inches(0.25), v_top + Inches(3.4),
           v_w - Inches(0.5), Inches(0.4),
           v["scope"], size=11, bold=True, color=PA_ORANGE_DARK,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # 底部最终主张（最高级别 takeaway — 最大字号）
    rect(s, Inches(0.5), Inches(6.45), Inches(12.33), Inches(0.75), PA_ORANGE_DEEP)
    tb(s, Inches(0.7), Inches(6.5), Inches(2.0), Inches(0.3),
       "FINAL VISION   ·   最终愿景",
       size=10, bold=True, color=PA_ORANGE_LT)
    tb(s, Inches(0.7), Inches(6.78), Inches(11.93), Inches(0.4),
       "客户在哪里 · 岗位 AI 跟到哪里 — 让平安在 AI 时代的综合金融客户经营上先发、跑通、领跑。",
       size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE,
       line_spacing=1.3)

    board_footer(s)


# ===== 生成 =====
slide_1()
slide_2()
slide_3()
slide_4()
slide_5()
slide_6()

out = "/projects/sandbox/kiro2026/银卡岗位1+N数字化转型_集团董事长汇报版.pptx"
prs.save(out)
print(f"已生成: {out}")
print(f"共 {len(prs.slides)} 页")
