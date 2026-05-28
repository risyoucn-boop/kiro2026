# -*- coding: utf-8 -*-
"""
小龙虾 AI 助手 · 新人上手培训课件（含封面共 10 页）
- 平安橙 + 白底
- 风格：优雅、克制、留白充足、培训感而非汇报感
- 受众：刚拿到 OpenClaw 账号的新员工
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ===== 平安橙 =====
PA_ORANGE      = RGBColor(0xEC, 0x6F, 0x1C)
PA_ORANGE_DARK = RGBColor(0xC8, 0x55, 0x0E)
PA_ORANGE_DEEP = RGBColor(0x7A, 0x32, 0x06)
PA_ORANGE_LT   = RGBColor(0xFF, 0xE8, 0xD2)
PA_ORANGE_BG   = RGBColor(0xFF, 0xF7, 0xEF)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK       = RGBColor(0x14, 0x14, 0x14)
DARK      = RGBColor(0x33, 0x33, 0x33)
GRAY      = RGBColor(0x6E, 0x6E, 0x6E)
MIDGRAY   = RGBColor(0xA0, 0xA0, 0xA0)
LIGHTGRAY = RGBColor(0xDD, 0xDD, 0xDD)
PALE      = RGBColor(0xF7, 0xF5, 0xF3)
CODE_BG   = RGBColor(0xFA, 0xF7, 0xF4)
CODE_BD   = RGBColor(0xE6, 0xCF, 0xB4)
WARN_BG   = RGBColor(0xFD, 0xEC, 0xE0)
RED_NO    = RGBColor(0xC6, 0x28, 0x28)

CN_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height


def set_run(run, text, size=11, bold=False, color=INK, italic=False, font=CN_FONT):
    run.text = text
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', CN_FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def tb(slide, l, t, w, h, text, size=11, bold=False, color=INK,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
       line_spacing=1.3, font=CN_FONT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        set_run(run, ln, size=size, bold=bold, color=color, italic=italic, font=font)
    return box


def rect(slide, l, t, w, h, fill, line=None, lw=0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def rounded(slide, l, t, w, h, fill, line=None, lw=0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.adjustments[0] = 0.15
    return s


def white_bg(slide):
    rect(slide, 0, 0, SW, SH, WHITE)


def page_header(slide, page, total, eyebrow, title, subtitle=None):
    rect(slide, 0, 0, SW, Inches(0.06), PA_ORANGE)
    tb(slide, Inches(0.5), Inches(0.22), Inches(8.0), Inches(0.28),
       eyebrow, size=10, bold=True, color=PA_ORANGE)
    tb(slide, Inches(11.4), Inches(0.22), Inches(1.4), Inches(0.28),
       f"{page} / {total}", size=10, bold=True, color=GRAY,
       align=PP_ALIGN.RIGHT)
    tb(slide, Inches(0.5), Inches(0.55), Inches(12.3), Inches(0.75),
       title, size=28, bold=True, color=INK,
       anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
    if subtitle:
        tb(slide, Inches(0.5), Inches(1.32), Inches(12.3), Inches(0.4),
           subtitle, size=13, color=GRAY, line_spacing=1.3)
        rule_y = Inches(1.85)
    else:
        rule_y = Inches(1.5)
    rect(slide, Inches(0.5), rule_y, Inches(0.5), Inches(0.04), PA_ORANGE)
    rect(slide, Inches(1.0), rule_y + Inches(0.014),
         Inches(11.83), Inches(0.012), LIGHTGRAY)


def page_footer(slide):
    rect(slide, Inches(0.5), Inches(7.18), Inches(12.33), Inches(0.012), LIGHTGRAY)
    tb(slide, Inches(0.5), Inches(7.22), Inches(8.0), Inches(0.22),
       "小龙虾新人上手课   ·   银卡客服中心",
       size=8.5, color=MIDGRAY)
    tb(slide, Inches(8.5), Inches(7.22), Inches(4.33), Inches(0.22),
       "POWERED BY OPENCLAW    ·    PING AN",
       size=8.5, bold=True, color=PA_ORANGE,
       align=PP_ALIGN.RIGHT)


def kicker_box(slide, l, t, w, h, label, body, fill=PA_ORANGE_BG,
               label_color=PA_ORANGE_DARK, body_color=INK,
               border=True):
    if border:
        rounded(slide, l, t, w, h, fill, line=PA_ORANGE, lw=1.0)
    else:
        rounded(slide, l, t, w, h, fill)
    rect(slide, l, t, Inches(0.08), h, PA_ORANGE)
    tb(slide, l + Inches(0.25), t + Inches(0.08),
       w - Inches(0.4), Inches(0.3),
       label, size=10, bold=True, color=label_color)
    tb(slide, l + Inches(0.25), t + Inches(0.42),
       w - Inches(0.4), h - Inches(0.5),
       body, size=13, bold=True, color=body_color, line_spacing=1.4)


def code_block(slide, l, t, w, h, lines, size=10):
    rounded(slide, l, t, w, h, CODE_BG, line=CODE_BD, lw=1.0)
    rect(slide, l + Inches(0.12), t + Inches(0.12), Inches(0.06), Inches(0.06), PA_ORANGE)
    rect(slide, l + Inches(0.22), t + Inches(0.12), Inches(0.06), Inches(0.06), PA_ORANGE_DARK)
    rect(slide, l + Inches(0.32), t + Inches(0.12), Inches(0.06), Inches(0.06), PA_ORANGE_DEEP)
    tb(slide, l + Inches(0.5), t + Inches(0.08),
       w - Inches(1.0), Inches(0.28),
       "PROMPT  提示词示例", size=8.5, bold=True, color=PA_ORANGE_DARK)
    tb(slide, l + Inches(0.3), t + Inches(0.4),
       w - Inches(0.6), h - Inches(0.5),
       lines, size=size, color=DARK, line_spacing=1.45,
       font=MONO_FONT)


# 封面
def slide_cover():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    rect(s, 0, 0, SW, Inches(0.08), PA_ORANGE)
    rect(s, 0, Inches(0.08), Inches(0.5), SH - Inches(0.08), PA_ORANGE_BG)
    tb(s, Inches(0.9), Inches(0.5), Inches(8.0), Inches(0.3),
       "OPENCLAW · 小龙虾 AI 助手", size=11, bold=True, color=PA_ORANGE)
    tb(s, Inches(11.4), Inches(0.5), Inches(1.4), Inches(0.3),
       "新人课件 · v1", size=10, bold=True, color=GRAY,
       align=PP_ALIGN.RIGHT)
    tb(s, Inches(0.9), Inches(1.3), Inches(11.5), Inches(0.4),
       "AI ONBOARDING   ·   新人 AI 助手第一堂上手课",
       size=12, bold=True, color=PA_ORANGE_DARK)
    tb(s, Inches(0.9), Inches(1.85), Inches(11.5), Inches(1.5),
       "认识小龙虾",
       size=72, bold=True, color=INK, line_spacing=1.0)
    tb(s, Inches(0.9), Inches(3.45), Inches(11.5), Inches(1.0),
       "从「会打开账号」 → 到「会用它完成工作」",
       size=24, bold=True, color=PA_ORANGE_DARK, line_spacing=1.3)
    rect(s, Inches(0.9), Inches(4.55), Inches(0.8), Inches(0.06), PA_ORANGE)
    tb(s, Inches(0.9), Inches(4.8), Inches(11.5), Inches(0.4),
       "学完后你能做到", size=11, bold=True, color=GRAY)
    goals = [
        ("01", "知道小龙虾是什么、能做什么、不能做什么"),
        ("02", "会用「角色 + 任务 + 背景 + 要求 + 输出格式」写提示词"),
        ("03", "能让小龙虾帮你写通知、整理纪要、拆解知识 — 今天就能用"),
    ]
    g_top = Inches(5.2)
    for i, (no, txt) in enumerate(goals):
        y = g_top + Inches(0.55) * i
        rect(s, Inches(0.9), y + Inches(0.08),
             Inches(0.42), Inches(0.42), PA_ORANGE)
        tb(s, Inches(0.9), y + Inches(0.08),
           Inches(0.42), Inches(0.42),
           no, size=12, bold=True, color=WHITE,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, Inches(1.45), y + Inches(0.1),
           Inches(11.0), Inches(0.45),
           txt, size=14, bold=True, color=INK,
           anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0, Inches(7.18), SW, Inches(0.012), LIGHTGRAY)
    tb(s, Inches(0.9), Inches(7.22), Inches(11.5), Inches(0.22),
       "银卡客服中心 · 数字化训练   ·   今天不教 AI 原理 — 今天教你怎么把这封邮件少写 20 分钟",
       size=10, italic=True, color=GRAY)


def slide_1():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    page_header(
        s, 1, 9,
        "WHY   ·   为什么要认识小龙虾",
        "小龙虾不是新工具 — 它是新的工作助手",
        "以后遇到「重复写、反复改、难下笔、没思路」的工作 — 让小龙虾先打一版底稿",
    )
    rect(s, Inches(0.5), Inches(2.1), Inches(6.0), Inches(0.5), PALE)
    tb(s, Inches(0.7), Inches(2.1), Inches(5.6), Inches(0.5),
       "▍ 原来   全部靠人手", size=12, bold=True, color=GRAY,
       anchor=MSO_ANCHOR.MIDDLE)
    old_items = [
        "写邮件 / 通知",  "整理会议纪要",
        "写培训小结",     "做材料初稿",
        "拆解业务知识",   "优化话术",
        "总结案例",       "生成表格框架",
    ]
    o_top = Inches(2.7)
    o_w = Inches(2.85)
    o_h = Inches(0.55)
    for i, t_ in enumerate(old_items):
        col = i % 2
        row = i // 2
        x = Inches(0.5) + (o_w + Inches(0.1)) * col
        y = o_top + (o_h + Inches(0.1)) * row
        rounded(s, x, y, o_w, o_h, WHITE, line=LIGHTGRAY, lw=0.75)
        rect(s, x + Inches(0.18), y + Inches(0.22),
             Inches(0.12), Inches(0.12), MIDGRAY)
        tb(s, x + Inches(0.4), y, o_w - Inches(0.5), o_h,
           t_, size=12, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(6.55), Inches(2.7), Inches(0.4), Inches(3.0),
       "→", size=44, bold=True, color=PA_ORANGE,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(7.0), Inches(2.1), Inches(5.85), Inches(0.5), PA_ORANGE)
    tb(s, Inches(7.2), Inches(2.1), Inches(5.6), Inches(0.5),
       "▍ 现在   小龙虾帮你做", size=12, bold=True, color=WHITE,
       anchor=MSO_ANCHOR.MIDDLE)
    new_verbs = [
        ("起草", "Draft"),     ("整理", "Organize"),
        ("拆解", "Break-down"), ("归纳", "Summarize"),
        ("优化", "Polish"),     ("改写", "Rewrite"),
        ("总结", "Conclude"),   ("生成", "Generate"),
    ]
    n_top = Inches(2.7)
    n_w = Inches(2.85)
    n_h = Inches(0.55)
    for i, (cn, en) in enumerate(new_verbs):
        col = i % 2
        row = i // 2
        x = Inches(7.0) + (n_w + Inches(0.1)) * col
        y = n_top + (n_h + Inches(0.1)) * row
        rounded(s, x, y, n_w, n_h, PA_ORANGE_BG, line=PA_ORANGE, lw=0.75)
        rect(s, x + Inches(0.15), y + Inches(0.18),
             Inches(0.18), Inches(0.18), PA_ORANGE)
        tb(s, x + Inches(0.45), y + Inches(0.06),
           Inches(1.4), Inches(0.25), cn,
           size=13, bold=True, color=INK)
        tb(s, x + Inches(0.45), y + Inches(0.3),
           Inches(2.0), Inches(0.22), en,
           size=8, color=PA_ORANGE_DARK, font="Calibri")
    kicker_box(
        s, Inches(0.5), Inches(6.4), Inches(12.33), Inches(0.7),
        "TAKEAWAY",
        "今天不是教大家一个复杂系统 — 而是让大家知道：动笔之前，先让小龙虾打一版底稿。",
    )
    page_footer(s)


def slide_2():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    page_header(
        s, 2, 9,
        "WHAT   ·   小龙虾是什么",
        "一个能帮你「想、写、改、整理」的 AI 工作助手",
        "不是搜索引擎，也不是简单问答机器人 — 它能根据你的任务和背景，生成、整理、拆解、优化",
    )
    rounded(s, Inches(0.5), Inches(2.1), Inches(12.33), Inches(0.65),
            PA_ORANGE_BG, line=PA_ORANGE, lw=1.0)
    rect(s, Inches(0.5), Inches(2.1), Inches(0.08), Inches(0.65), PA_ORANGE)
    tb(s, Inches(0.75), Inches(2.18), Inches(12.0), Inches(0.5),
       "小龙虾  =  基于大模型能力的 AI 工作助手   ·   能根据你的任务与背景，完成文字生成、内容整理、知识拆解、方案设计、话术优化、材料润色等工作。",
       size=12, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE,
       line_spacing=1.3)
    tb(s, Inches(0.5), Inches(2.95), Inches(7.5), Inches(0.32),
       "▍ 擅长什么   6 类能力", size=12, bold=True, color=PA_ORANGE_DARK)
    skills = [
        ("写", "邮件 · 通知 · 总结 · 汇报"),
        ("改", "改语气 · 改结构 · 改表达"),
        ("拆", "拆制度 · 拆流程 · 拆知识点"),
        ("总", "总结会议 · 材料 · 案例"),
        ("想", "提供思路 · 框架 · 方案"),
        ("练", "生成题目 · 话术 · 演练脚本"),
    ]
    sk_top = Inches(3.4)
    sk_w = Inches(2.45)
    sk_h = Inches(1.35)
    for i, (k, v) in enumerate(skills):
        col = i % 3
        row = i // 3
        x = Inches(0.5) + (sk_w + Inches(0.08)) * col
        y = sk_top + (sk_h + Inches(0.12)) * row
        rounded(s, x, y, sk_w, sk_h, WHITE, line=PA_ORANGE_LT, lw=1.0)
        rect(s, x, y, Inches(0.85), sk_h, PA_ORANGE)
        tb(s, x, y, Inches(0.85), sk_h,
           k, size=36, bold=True, color=WHITE,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + Inches(1.0), y + Inches(0.18),
           sk_w - Inches(1.1), Inches(1.0),
           v, size=10.5, color=DARK, line_spacing=1.5,
           anchor=MSO_ANCHOR.MIDDLE)
    rt_x = Inches(8.55)
    rt_w = Inches(4.28)
    tb(s, rt_x, Inches(2.95), rt_w, Inches(0.32),
       "▍ 不适合做什么   边界提醒", size=12, bold=True, color=RED_NO)
    nots = [
        ("直接替代业务判断",   "最终责任还是人"),
        ("处理敏感客户信息",   "存在信息安全风险"),
        ("直接生成对外承诺",   "可能口径不准"),
        ("不校验就复制使用",   "可能有错误或不贴合业务"),
    ]
    n_top = Inches(3.4)
    n_h = Inches(0.7)
    for i, (k, v) in enumerate(nots):
        y = n_top + (n_h + Inches(0.05)) * i
        rounded(s, rt_x, y, rt_w, n_h, WARN_BG, line=RED_NO, lw=0.75)
        rect(s, rt_x, y, Inches(0.08), n_h, RED_NO)
        tb(s, rt_x + Inches(0.25), y + Inches(0.05),
           rt_w - Inches(0.4), Inches(0.3),
           f"✕  {k}", size=11, bold=True, color=RED_NO)
        tb(s, rt_x + Inches(0.25), y + Inches(0.36),
           rt_w - Inches(0.4), Inches(0.3),
           v, size=10, color=DARK)
    kicker_box(
        s, Inches(0.5), Inches(6.4), Inches(12.33), Inches(0.7),
        "记住一句话",
        "小龙虾不是来替你工作 — 它是来帮你把「重复、零散、难下笔」的工作变快。",
    )
    page_footer(s)


def slide_3():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    page_header(
        s, 3, 9,
        "VOCABULARY   ·   关键词汇",
        "几个 AI 词汇 — 用人话讲清楚",
        "不讲深、不讲全 — 只讲「够用版」，新人记住这 7 个词就够了",
    )
    headers = ["词汇", "用人话讲", "和我们的关系"]
    rows = [
        ["大模型",        "见过大量文字和知识的 AI 系统",       "小龙虾背后的能力来源"],
        ["生成式 AI",     "能根据要求生成内容的 AI",            "可以帮我们写材料、写话术"],
        ["Prompt 提示词", "你给 AI 的任务说明",                 "说得越清楚 · 结果越好"],
        ["上下文",        "你提供的背景信息",                    "背景越完整 · 回答越贴近实际"],
        ["多轮对话",      "可以一直追问，让结果越改越好",        "不要指望一次问完"],
        ["幻觉",          "AI 一本正经地说错话",                 "业务内容必须人工校验"],
        ["工作流",        "把一个任务拆成多个步骤",              "让小龙虾产出更稳定"],
    ]
    n_rows = len(rows) + 1
    n_cols = len(headers)
    shp = s.shapes.add_table(n_rows, n_cols,
                              Inches(0.5), Inches(2.1),
                              Inches(12.33), Inches(3.45))
    tbl = shp.table
    cw = [2.3, 4.5, 5.5]
    s_ = sum(cw)
    for i, c in enumerate(cw):
        tbl.columns[i].width = int(Inches(12.33) * c / s_)
    for ci, hd in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PA_ORANGE
        cell.margin_left = cell.margin_right = Inches(0.15)
        cell.margin_top = cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tfx = cell.text_frame
        p = tfx.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        set_run(run, hd, size=12, bold=True, color=WHITE)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri + 1, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 == 0 else PA_ORANGE_BG
            cell.margin_left = cell.margin_right = Inches(0.15)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tfx = cell.text_frame
            p = tfx.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            if ci == 0:
                set_run(run, val, size=13, bold=True, color=PA_ORANGE_DARK)
            else:
                set_run(run, val, size=11, color=DARK)
    kicker_box(
        s, Inches(0.5), Inches(5.7), Inches(12.33), Inches(1.4),
        "金句   ·   记住这一句就够",
        "AI 不是「问一句、它就完美回答」 — AI 更像一个新人助手：你要交代清楚背景、目标、标准和输出格式。",
    )
    page_footer(s)


def slide_4():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    page_header(
        s, 4, 9,
        "PROMPT ENGINEERING   ·   提示词工程",
        "提示词不是「咒语」 — 是「任务说明书」",
        "用不好小龙虾，多数时候不是工具不行 — 是问的人没说清楚",
    )
    err_x = Inches(0.5)
    err_w = Inches(6.05)
    rounded(s, err_x, Inches(2.1), err_w, Inches(4.3),
            WHITE, line=RED_NO, lw=1.25)
    rect(s, err_x, Inches(2.1), err_w, Inches(0.55), WARN_BG)
    tb(s, err_x + Inches(0.25), Inches(2.1),
       err_w - Inches(0.4), Inches(0.55),
       "✕  错误问法   太短、太模糊、AI 猜不到",
       size=13, bold=True, color=RED_NO, anchor=MSO_ANCHOR.MIDDLE)
    code_block(s, err_x + Inches(0.25), Inches(2.85),
               err_w - Inches(0.5), Inches(0.65),
               ["帮我写个培训通知。"], size=14)
    tb(s, err_x + Inches(0.25), Inches(3.65),
       err_w - Inches(0.5), Inches(0.32),
       "AI 完全不知道：", size=11, bold=True, color=RED_NO)
    miss = [
        "•  给谁看？",
        "•  什么培训？",
        "•  语气正式还是轻松？",
        "•  要不要写参训要求？",
        "•  有没有训后作业？",
        "•  输出邮件还是微信？",
    ]
    tb(s, err_x + Inches(0.5), Inches(4.0),
       err_w - Inches(0.7), Inches(2.2),
       miss, size=12, color=DARK, line_spacing=1.5)
    tb(s, Inches(6.55), Inches(2.1), Inches(0.4), Inches(4.3),
       "→", size=44, bold=True, color=PA_ORANGE,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ok_x = Inches(7.0)
    ok_w = Inches(5.83)
    rounded(s, ok_x, Inches(2.1), ok_w, Inches(4.3),
            WHITE, line=PA_ORANGE, lw=1.25)
    rect(s, ok_x, Inches(2.1), ok_w, Inches(0.55), PA_ORANGE)
    tb(s, ok_x + Inches(0.25), Inches(2.1),
       ok_w - Inches(0.4), Inches(0.55),
       "✓  正确问法   把任务交代清楚",
       size=13, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    correct_lines = [
        "你是一名企业培训专员。",
        "请帮我写一封培训通知邮件，培训主题是",
        "「小龙虾基础使用培训」，对象是新开通",
        "OpenClaw 账号的员工。",
        "",
        "邮件要求：正式、简洁，包含 ——",
        "  · 培训背景    · 培训目标    · 培训内容",
        "  · 参训要求    · 训后作业",
        "",
        "输出格式:邮件标题 + 正文。",
    ]
    code_block(s, ok_x + Inches(0.25), Inches(2.85),
               ok_w - Inches(0.5), Inches(3.45),
               correct_lines, size=11)
    kicker_box(
        s, Inches(0.5), Inches(6.55), Inches(12.33), Inches(0.55),
        "TAKEAWAY",
        "提示词写得好 — 本质上就是把任务交代清楚。说不清楚，AI 永远做不好。",
    )
    page_footer(s)


def slide_5():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    page_header(
        s, 5, 9,
        "FRAMEWORK   ·   万能提示词结构",
        "一个好提示词，至少讲清楚 5 件事",
        "新人先不要追求高级技巧 — 把这 5 件事讲清楚，结果就明显变好",
    )
    steps = [
        ("01", "角色",      "ROLE",     "让小龙虾站在谁的角度",  "你是一名企业内训师"),
        ("02", "任务",      "TASK",     "要它完成什么",          "请设计一份培训通知"),
        ("03", "背景",      "CONTEXT",  "为什么要做",            "新员工刚拿到 OpenClaw 账号"),
        ("04", "要求",      "RULES",    "有什么限制",            "正式 · 不要太长 · 可直接发送"),
        ("05", "输出格式",  "FORMAT",   "按什么形式交付",        "邮件标题 + 正文"),
    ]
    s_top = Inches(2.1)
    s_h = Inches(2.0)
    s_w = Inches(2.42)
    for i, (no, name, en, desc, ex) in enumerate(steps):
        x = Inches(0.5) + (s_w + Inches(0.05)) * i
        rounded(s, x, s_top, s_w, s_h, WHITE, line=LIGHTGRAY, lw=0.75)
        rect(s, x, s_top, s_w, Inches(0.55), PA_ORANGE)
        tb(s, x + Inches(0.2), s_top + Inches(0.06),
           Inches(0.6), Inches(0.43),
           no, size=14, bold=True, color=PA_ORANGE_LT,
           anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + Inches(0.85), s_top, s_w - Inches(1.0), Inches(0.55),
           name, size=18, bold=True, color=WHITE,
           anchor=MSO_ANCHOR.MIDDLE)
        tb(s, x + Inches(0.2), s_top + Inches(0.65),
           s_w - Inches(0.3), Inches(0.25),
           en, size=9, bold=True, color=PA_ORANGE,
           font="Calibri")
        tb(s, x + Inches(0.2), s_top + Inches(0.92),
           s_w - Inches(0.3), Inches(0.45),
           desc, size=10.5, color=DARK, line_spacing=1.4)
        rounded(s, x + Inches(0.2), s_top + Inches(1.45),
                s_w - Inches(0.3), Inches(0.5),
                PA_ORANGE_BG)
        tb(s, x + Inches(0.3), s_top + Inches(1.48),
           s_w - Inches(0.5), Inches(0.18),
           "示例", size=8, bold=True, color=GRAY)
        tb(s, x + Inches(0.3), s_top + Inches(1.65),
           s_w - Inches(0.5), Inches(0.3),
           ex, size=10, bold=True, color=PA_ORANGE_DARK,
           anchor=MSO_ANCHOR.MIDDLE)
    tb(s, Inches(0.5), Inches(4.3), Inches(12.33), Inches(0.32),
       "▍ 新人可复制模板   把【】里的内容换成你的真实任务",
       size=12, bold=True, color=PA_ORANGE_DARK)
    template = [
        "你是一名【角色】。",
        "请围绕【具体任务】帮我完成【输出成果】。",
        "背景是：【补充业务背景、对象、使用场景】",
        "要求：1. 【要求1】   2. 【要求2】   3. 【要求3】",
        "请按照以下格式输出：一、【模块1】   二、【模块2】   三、【模块3】",
    ]
    code_block(s, Inches(0.5), Inches(4.65),
               Inches(12.33), Inches(1.7),
               template, size=11)
    kicker_box(
        s, Inches(0.5), Inches(6.45), Inches(12.33), Inches(0.65),
        "TAKEAWAY",
        "角色 + 任务 + 背景 + 要求 + 输出格式  ——  五件事讲清楚 · 结果立刻变好。",
    )
    page_footer(s)


def render_practice(s, page, total, eyebrow, title, scene_desc,
                     pain_points, prompt_lines, value_lines):
    page_header(s, page, total, eyebrow, title,
                "今天学完，今天就能用 — 一个真实工作场景，一段可复制提示词")
    sc_x = Inches(0.5)
    sc_w = Inches(5.5)
    sc_top = Inches(2.1)
    tb(s, sc_x, sc_top, sc_w, Inches(0.32),
       "▍ 场景",
       size=12, bold=True, color=PA_ORANGE_DARK)
    rounded(s, sc_x, sc_top + Inches(0.4), sc_w, Inches(0.85),
            PA_ORANGE_BG, line=PA_ORANGE_LT, lw=1.0)
    tb(s, sc_x + Inches(0.25), sc_top + Inches(0.45),
       sc_w - Inches(0.5), Inches(0.78),
       scene_desc, size=12, bold=True, color=INK,
       anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.4)
    tb(s, sc_x, sc_top + Inches(1.4), sc_w, Inches(0.32),
       "▍ 原始痛点",
       size=12, bold=True, color=GRAY)
    pain_h_each = Inches(0.42)
    for i, p in enumerate(pain_points):
        y = sc_top + Inches(1.78) + pain_h_each * i
        rect(s, sc_x, y + Inches(0.15), Inches(0.1), Inches(0.1), MIDGRAY)
        tb(s, sc_x + Inches(0.25), y, sc_w - Inches(0.3), pain_h_each,
           p, size=11.5, color=DARK, anchor=MSO_ANCHOR.MIDDLE)
    pr_x = Inches(6.25)
    pr_w = Inches(6.58)
    tb(s, pr_x, sc_top, pr_w, Inches(0.32),
       "▍ 提示词示例   可直接复制",
       size=12, bold=True, color=PA_ORANGE_DARK)
    code_block(s, pr_x, sc_top + Inches(0.4),
               pr_w, Inches(4.55),
               prompt_lines, size=10.5)
    val_top = Inches(6.4)
    rounded(s, Inches(0.5), val_top, Inches(12.33), Inches(0.7),
            PA_ORANGE, line=PA_ORANGE_DEEP, lw=1.0)
    tb(s, Inches(0.7), val_top + Inches(0.08),
       Inches(2.0), Inches(0.3),
       "VALUE   ·   它给你节省了什么",
       size=10, bold=True, color=PA_ORANGE_LT)
    tb(s, Inches(0.7), val_top + Inches(0.36),
       Inches(11.93), Inches(0.32),
       value_lines, size=12.5, bold=True, color=WHITE,
       line_spacing=1.3)
    page_footer(s)


def slide_6():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    prompt = [
        "你是一名企业培训专员。",
        "请帮我写一封培训通知邮件，主题是「小龙虾基础使用培训」。",
        "",
        "背景：",
        "公司已为部分新员工开通 OpenClaw 账号，希望通过本次",
        "培训帮助大家了解小龙虾的基础用法、典型场景和注意事项。",
        "",
        "要求：",
        "  1. 语气正式、清晰、适合企业内部邮件",
        "  2. 包含培训背景、培训对象、培训内容、参训要求、训后作业",
        "  3. 不要空泛，不要写得太长",
        "  4. 可直接复制发送",
        "",
        "请输出：邮件标题 + 邮件正文",
    ]
    render_practice(
        s, 6, 9,
        "PRACTICE 1   ·   实战一",
        "让小龙虾帮你写一封培训通知",
        "培训岗要通知新员工参加小龙虾使用培训。",
        ["自己想标题、写背景、列内容",
         "反复改语气、来回调整",
         "容易漏写参训要求 / 训后作业",
         "一封邮件半小时甚至更久"],
        prompt,
        "从「自己从零开始写」 → 变成「先生成初稿、再人工校正」 — 一封邮件最快 5 分钟搞定。",
    )


def slide_7():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    prompt = [
        "你是一名企业会议纪要整理专家。",
        "请帮我把下面这段会议记录整理成结构清晰的会议纪要。",
        "",
        "要求：",
        "  1. 保留真实意思，不要编造",
        "  2. 口语化内容改成正式表达",
        "  3. 提炼会议重点，整理待办事项",
        "  4. 标注责任人、时间节点；原文无则写「未明确」",
        "  5. 不确定的地方请标注「待确认」",
        "",
        "输出格式：",
        "一、会议主题   二、会议重点   三、关键决策",
        "四、待办事项   五、风险提醒",
        "",
        "原始记录：【粘贴会议内容】",
    ]
    render_practice(
        s, 7, 9,
        "PRACTICE 2   ·   实战二",
        "把混乱记录整理成会议纪要",
        "开完会后，有一段比较乱的会议记录，需要整理成正式纪要。",
        ["记录口语化、不成型",
         "重点不清楚、待办不明确",
         "责任人和时间节点容易漏",
         "整理一份纪要要 1 ~ 2 小时"],
        prompt,
        "小龙虾负责「整理第一稿」 — 你只需要做「最终人工确认」 — 时间从 2 小时缩到 20 分钟。",
    )


def slide_8():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    prompt = [
        "你是一名企业业务培训师。",
        "请帮我把下面这段业务知识拆解成新人容易理解的培训材料。",
        "",
        "要求：",
        "  1. 先用一句话说明这个知识点解决什么问题",
        "  2. 提炼核心规则",
        "  3. 说明适用场景",
        "  4. 列出客户可能会问的问题",
        "  5. 给出员工可参考的话术",
        "  6. 标注容易出错的地方",
        "",
        "输出格式：",
        "一、一句话理解   二、核心规则   三、适用场景",
        "四、客户常见问法 五、员工参考话术 六、易错提醒",
        "",
        "原始内容：【粘贴业务知识】",
    ]
    render_practice(
        s, 8, 9,
        "PRACTICE 3   ·   实战三",
        "把复杂知识拆成新人能看懂的内容",
        "新人看到制度、流程、业务规则，经常觉得文件长、抓不住重点。",
        ["文件太长、密密麻麻",
         "重点不清楚、规则模糊",
         "不知道客户会怎么问",
         "不知道实际工作怎么用"],
        prompt,
        "把「长材料」 → 变成「可学习、可培训、可使用」的结构化内容 — 新人上手周期缩一半。",
    )


def slide_9():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    white_bg(s)
    page_header(
        s, 9, 9,
        "WRAP UP   ·   课程收尾",
        "小龙虾使用的三个关键习惯",
        "学完不一定马上做复杂项目 — 但请把这三个习惯，从今天开始养成",
    )
    habits = [
        {
            "no": "01",
            "title": "先说清楚任务",
            "en": "BE SPECIFIC",
            "body": "不要只说「帮我写一下」 — \n说清楚对象、背景、目标和格式。",
            "tip": "记忆口诀：角色 + 任务 + 背景 + 要求 + 输出格式",
        },
        {
            "no": "02",
            "title": "把 AI 当初稿助手",
            "en": "AI DRAFTS · YOU JUDGE",
            "body": "AI 负责起草和整理 — \n你负责判断、把关、最终交付。",
            "tip": "业务内容必须人工校验 — AI 偶尔会一本正经地说错",
        },
        {
            "no": "03",
            "title": "把好用的方法存下来",
            "en": "BUILD YOUR LIBRARY",
            "body": "每次用得好的提示词 — \n沉淀成模板，下次直接复用。",
            "tip": "三个月后，你会有一本自己的「提示词工具箱」",
        },
    ]
    h_top = Inches(2.15)
    h_w = Inches(4.05)
    h_h = Inches(3.95)
    for i, h in enumerate(habits):
        x = Inches(0.5) + (h_w + Inches(0.13)) * i
        rounded(s, x, h_top, h_w, h_h, WHITE, line=LIGHTGRAY, lw=0.75)
        rect(s, x, h_top, h_w, Inches(0.85), PA_ORANGE)
        tb(s, x + Inches(0.25), h_top + Inches(0.1),
           Inches(0.8), Inches(0.4),
           h["no"], size=20, bold=True, color=PA_ORANGE_LT,
           italic=True)
        tb(s, x + Inches(0.25), h_top + Inches(0.45),
           h_w - Inches(0.4), Inches(0.35),
           h["en"], size=10, bold=True, color=PA_ORANGE_LT)
        tb(s, x + Inches(0.3), h_top + Inches(1.0),
           h_w - Inches(0.5), Inches(0.55),
           h["title"], size=22, bold=True, color=INK)
        rect(s, x + Inches(0.3), h_top + Inches(1.65),
             Inches(0.5), Inches(0.04), PA_ORANGE)
        tb(s, x + Inches(0.3), h_top + Inches(1.85),
           h_w - Inches(0.5), Inches(1.2),
           h["body"], size=12, color=DARK, line_spacing=1.55)
        rounded(s, x + Inches(0.3), h_top + Inches(3.05),
                h_w - Inches(0.5), Inches(0.78),
                PA_ORANGE_BG)
        tb(s, x + Inches(0.4), h_top + Inches(3.1),
           h_w - Inches(0.7), Inches(0.25),
           "TIP", size=8, bold=True, color=PA_ORANGE_DARK)
        tb(s, x + Inches(0.4), h_top + Inches(3.32),
           h_w - Inches(0.7), Inches(0.5),
           h["tip"], size=10, bold=True, color=PA_ORANGE_DARK,
           line_spacing=1.4)
    rect(s, Inches(0.5), Inches(6.3), Inches(12.33), Inches(0.85),
         PA_ORANGE_DEEP)
    tb(s, Inches(0.7), Inches(6.38),
       Inches(2.5), Inches(0.3),
       "FINAL TAKEAWAY   ·   结尾金句",
       size=10, bold=True, color=PA_ORANGE_LT)
    tb(s, Inches(0.7), Inches(6.65),
       Inches(11.93), Inches(0.5),
       "小龙虾真正的价值 — 不是替我们工作，而是帮我们把「重复、零散、难下笔」的工作，更快变成可交付的成果。",
       size=14, bold=True, color=WHITE,
       anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.3)
    page_footer(s)


slide_cover()
slide_1()
slide_2()
slide_3()
slide_4()
slide_5()
slide_6()
slide_7()
slide_8()
slide_9()

out = "/projects/sandbox/kiro2026/小龙虾AI助手_新人上手课件.pptx"
prs.save(out)
print(f"已生成: {out}")
print(f"共 {len(prs.slides)} 页（含封面）")
