"""
Scrape emojis from https://www.emojiall.com/zh-hant and produce an Excel
file with columns A-G:
  A: 序号 (sequence)
  B: Emoji
  C: 中文名 (Chinese short name)
  D: 场景章节 (Category / Subcategory from the website)
  E: 建议用法 (Suggested usage — first paragraph of 意义与描述)
  F: 关键词 (Keywords, comma-separated)
  G: Code Point (e.g. U+1F4CA)
"""

import concurrent.futures
import json
import os
import re
import sys
import time
import urllib.parse
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "https://www.emojiall.com"
LANG_PATH = "/zh-hant"
CATEGORIES = list("ABCDEFGHIJ")  # the 10 top-level categories
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "zh-Hant,zh;q=0.9,en;q=0.8",
}
OUT_DIR = Path(__file__).resolve().parent.parent / "data"
OUT_DIR.mkdir(parents=True, exist_ok=True)
LISTING_JSON = OUT_DIR / "listing.json"
DETAILS_JSON = OUT_DIR / "details.json"
XLSX_PATH = Path(__file__).resolve().parent.parent / "emojiall_zh_hant_emojis.xlsx"

session = requests.Session()
session.headers.update(HEADERS)


def get(url: str, retries: int = 3, sleep: float = 1.0) -> str:
    last_exc = None
    for i in range(retries):
        try:
            r = session.get(url, timeout=30)
            if r.status_code == 200:
                return r.text
            last_exc = RuntimeError(f"HTTP {r.status_code} for {url}")
        except Exception as e:
            last_exc = e
        time.sleep(sleep * (i + 1))
    raise last_exc  # type: ignore[misc]


# ---------- Step 1: collect emoji listings from category pages ----------

CAT_NAME_RE = re.compile(r"類別:([^\u3000\s]+?)相關表情符號合集")


def extract_category_name(soup: BeautifulSoup) -> str:
    # Try og:title meta
    m = soup.find("meta", attrs={"property": "og:title"})
    if m and m.get("content"):
        match = CAT_NAME_RE.search(m["content"])
        if match:
            name = match.group(1)
            # strip leading emoji chars (non-CJK before CJK), keep the Chinese name
            cleaned = re.sub(r"^[^\u4e00-\u9fff]+", "", name)
            return cleaned or name
    return ""


def parse_subsection_title(text: str) -> str:
    """Subsection h2 looks like '😄笑臉14' — strip emoji + trailing digits."""
    s = re.sub(r"\d+$", "", text).strip()
    # Strip leading non-Han characters (emoji prefix)
    s = re.sub(r"^[^\u4e00-\u9fff]+", "", s)
    return s.strip()


def scrape_category(cat_letter: str) -> list[dict]:
    url = f"{BASE}{LANG_PATH}/categories/{cat_letter}"
    print(f"[cat] fetching {url}", flush=True)
    html = get(url)
    soup = BeautifulSoup(html, "lxml")
    cat_name = extract_category_name(soup) or cat_letter

    # Iterate top-level h2 subsections inside the main content. Each subsection
    # heading is followed by a block containing emoji <a> links.
    items: list[dict] = []
    seen: set[str] = set()
    # Find h2 headings that contain at least one emoji link in following siblings
    for h2 in soup.find_all("h2"):
        sub_title_raw = h2.get_text(strip=True)
        # Skip non-content headings (导航, 语言, etc.)
        sub_title = parse_subsection_title(sub_title_raw)
        # Walk siblings until we hit the next h2
        emoji_links: list = []
        for sib in h2.find_next_siblings():
            if sib.name == "h2":
                break
            for a in sib.find_all("a", href=True):
                if a["href"].startswith(f"{LANG_PATH}/emoji/"):
                    emoji_links.append(a)
        if not emoji_links:
            continue
        # Each emoji is represented by 2 <a> tags: the emoji char then the name.
        # Iterate pairs.
        i = 0
        while i < len(emoji_links):
            a1 = emoji_links[i]
            href = a1["href"]
            text1 = a1.get_text(strip=True)
            name = ""
            # If next link has same href and a text label, use it as the name
            if i + 1 < len(emoji_links) and emoji_links[i + 1]["href"] == href:
                name = emoji_links[i + 1].get_text(strip=True)
                i += 2
            else:
                # fallback to title attr
                name = a1.get("title", "") or ""
                i += 1
            if href in seen:
                continue
            seen.add(href)
            # Decode the URL-encoded emoji char
            try:
                emoji_char = urllib.parse.unquote(href.split("/zh-hant/emoji/")[-1])
            except Exception:
                emoji_char = text1
            # Skip if 'emoji char' text is empty
            if not emoji_char:
                continue
            items.append({
                "category": cat_name,
                "subcategory": sub_title,
                "name_listing": name,
                "emoji": emoji_char,
                "url": BASE + href,
            })
    print(f"[cat] {cat_letter} ({cat_name}): {len(items)} unique emojis", flush=True)
    return items


def scrape_listings() -> list[dict]:
    if LISTING_JSON.exists():
        print(f"[cache] loading {LISTING_JSON}")
        return json.loads(LISTING_JSON.read_text(encoding="utf-8"))
    all_items: list[dict] = []
    seen_urls: set[str] = set()
    for c in CATEGORIES:
        items = scrape_category(c)
        for it in items:
            if it["url"] not in seen_urls:
                seen_urls.add(it["url"])
                all_items.append(it)
        time.sleep(0.3)
    LISTING_JSON.write_text(
        json.dumps(all_items, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[cache] saved {len(all_items)} items -> {LISTING_JSON}")
    return all_items


# ---------- Step 2: fetch emoji detail pages ----------


def parse_emoji_page(html: str) -> dict:
    soup = BeautifulSoup(html, "lxml")
    info: dict = {}
    # Find the basic info table by looking for 'Emoji:' cell or '字符編碼'
    for tbl in soup.find_all("table"):
        rows = tbl.find_all("tr")
        if not rows:
            continue
        first_text = rows[0].get_text(" ", strip=True)
        # Pick the table whose first row contains 'Emoji'
        if "Emoji" not in first_text and not any(
            "字符編碼" in r.get_text() for r in rows[:5]
        ):
            continue
        for tr in rows:
            cells = tr.find_all(["th", "td"])
            if len(cells) < 2:
                continue
            key = cells[0].get_text(" ", strip=True).rstrip(":：")
            val = cells[1].get_text(" ", strip=True)
            if key:
                info[key] = val
        break

    # Description / 意义与描述
    desc = ""
    heading = soup.find(
        lambda t: t.name == "h2" and "意義與描述" in t.get_text()
    )
    if heading:
        for sib in heading.find_next_siblings():
            if sib.name == "h2":
                break
            t = sib.get_text(" ", strip=True)
            if t:
                desc = t
                break
    info["_description"] = desc
    return info


def fetch_detail(item: dict) -> dict:
    try:
        html = get(item["url"])
        parsed = parse_emoji_page(html)
        return {"url": item["url"], "ok": True, "data": parsed}
    except Exception as e:
        return {"url": item["url"], "ok": False, "error": str(e)}


def scrape_details(items: list[dict]) -> dict:
    cache: dict = {}
    if DETAILS_JSON.exists():
        cache = json.loads(DETAILS_JSON.read_text(encoding="utf-8"))
        print(f"[cache] loaded {len(cache)} details from disk")
    todo = [it for it in items if it["url"] not in cache]
    print(f"[detail] {len(todo)}/{len(items)} pages to fetch")
    if not todo:
        return cache
    save_every = 100
    counter = 0
    with concurrent.futures.ThreadPoolExecutor(max_workers=20) as ex:
        futs = {ex.submit(fetch_detail, it): it for it in todo}
        for f in concurrent.futures.as_completed(futs):
            it = futs[f]
            res = f.result()
            cache[it["url"]] = res
            counter += 1
            if counter % 50 == 0:
                print(f"[detail] {counter}/{len(todo)}", flush=True)
            if counter % save_every == 0:
                DETAILS_JSON.write_text(
                    json.dumps(cache, ensure_ascii=False, indent=1),
                    encoding="utf-8",
                )
    DETAILS_JSON.write_text(
        json.dumps(cache, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    return cache


# ---------- Step 3: build Excel ----------


def normalize_keywords(s: str) -> str:
    if not s:
        return ""
    # Site uses '|' as separator; turn into Chinese comma '、' to match user sample
    parts = [p.strip() for p in re.split(r"[|｜,，]", s) if p.strip()]
    return "、".join(parts)


def codepoint_from_emoji(emoji: str) -> str:
    """Compute 'U+XXXX [U+XXXX...]' directly from the emoji character itself.

    Matches the style in the reference image: each codepoint prefixed with U+,
    space-separated. ZWJ (U+200D) is included as it's a real codepoint.
    """
    if not emoji:
        return ""
    parts = [f"U+{ord(ch):04X}" for ch in emoji]
    return " ".join(parts)


def normalize_codepoint(s: str) -> str:
    if not s:
        return ""
    # Strip trailing 复制 / 複製 button text
    s = s.replace("複製", "").replace("复制", "").strip()
    # Some pages include trailing per-codepoint annotations after the main
    # codepoint sequence, e.g. "U+1F590 FE0F  1F590 - 🖐 停止 FE0F - ️ 變體選擇符-16".
    # Cut everything from the first " - " separator.
    sep_idx = s.find(" - ")
    if sep_idx > 0:
        s = s[:sep_idx]
    # Extract just the U+XXXX [XXXX [XXXX...]] portion at the start.
    m = re.match(r"^(U\+\s*[0-9A-Fa-f]+(?:\s+[0-9A-Fa-f]+)*)", s)
    if m:
        s = m.group(1)
    # Remove the space right after "U+" and collapse runs of spaces.
    s = re.sub(r"^U\+\s+", "U+", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def normalize_name(zh: str, fallback: str = "") -> str:
    return (zh or fallback or "").strip()


def normalize_description(s: str) -> str:
    if not s:
        return ""
    # Site descriptions sometimes contain trailing 'Xuechao 💡 拓展閱讀與科普 ...' boilerplate.
    # Cut at the first such boilerplate marker.
    cut_markers = [
        "💡 拓展閱讀與科普",
        "拓展閱讀與科普",
        "💡拓展閱讀與科普",
    ]
    for m in cut_markers:
        idx = s.find(m)
        if idx >= 0:
            s = s[:idx]
            break
    # Trim trailing author name patterns (single English token followed by 💡 cut already)
    s = re.sub(r"\s*[A-Za-z][A-Za-z0-9_]*\s*$", "", s).strip()
    # Compact internal whitespace
    s = re.sub(r"\s+", " ", s).strip()
    # Truncate overly long descriptions (some pages include long quoted ads/poems).
    # Keep first sentence cluster up to a soft cap.
    MAX = 280
    if len(s) > MAX:
        # Try cutting at the nearest Chinese full-stop/period after the cap.
        cut = -1
        for ch in ("。", "！", "？", "."):
            i = s.find(ch, MAX // 2, MAX + 60)
            if i >= 0 and (cut < 0 or i < cut):
                cut = i
        if cut > 0:
            s = s[: cut + 1]
        else:
            s = s[:MAX].rstrip() + "…"
    return s


def build_excel(items: list[dict], details: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Emojis"
    headers = ["序号", "Emoji", "中文名", "场景章节", "建议用法", "关键词", "Code Point"]
    ws.append(headers)

    # Header style
    header_fill = PatternFill("solid", fgColor="FFE5B4")
    header_font = Font(bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for i, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center

    seq = 0
    for it in items:
        url = it["url"]
        d = details.get(url, {})
        data = (d or {}).get("data", {}) if d.get("ok") else {}
        zh_name = normalize_name(data.get("簡短名稱", ""), it.get("name_listing", ""))
        # Always compute code point from the emoji char itself for consistency
        # (the site sometimes interleaves per-codepoint annotations into the
        # 字符編碼 cell, which is hard to clean reliably).
        code_point = codepoint_from_emoji(it["emoji"])
        keywords = normalize_keywords(data.get("關鍵字", ""))
        desc = normalize_description(data.get("_description", ""))
        # 場景章節 = Category / Subcategory (subcategory if available)
        cat = it.get("category", "")
        sub = it.get("subcategory", "")
        scene = f"{cat} / {sub}" if sub else cat

        seq += 1
        ws.append([
            seq,
            it["emoji"],
            zh_name,
            scene,
            desc,
            keywords,
            code_point,
        ])

    # Style data rows
    for r in range(2, ws.max_row + 1):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            if col in (1, 2, 7):
                cell.alignment = center
            else:
                cell.alignment = left_wrap

    # Column widths approximating the user's sample
    widths = {
        "A": 8,    # 序号
        "B": 8,    # Emoji
        "C": 22,   # 中文名
        "D": 22,   # 场景章节
        "E": 60,   # 建议用法
        "F": 50,   # 关键词
        "G": 22,   # Code Point
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_PATH)
    print(f"[xlsx] wrote {XLSX_PATH} with {seq} rows")
    return XLSX_PATH


def main():
    items = scrape_listings()
    print(f"Total unique emoji listings: {len(items)}")
    details = scrape_details(items)
    ok = sum(1 for v in details.values() if v.get("ok"))
    print(f"Detail pages parsed OK: {ok}/{len(details)}")
    build_excel(items, details)


if __name__ == "__main__":
    main()
